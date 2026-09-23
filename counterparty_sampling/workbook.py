from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .core import AlertSamplingResult, SamplingError, Transaction, sample_alert


def _text(value: object) -> str:
    return "" if value is None else " ".join(str(value).strip().split())


def _header_map(headers: Iterable[object]) -> dict[str, int]:
    return {_text(header).casefold(): index for index, header in enumerate(headers)}


def _require_headers(headers: dict[str, int], required: Iterable[str], source: str) -> None:
    missing = [name for name in required if name.casefold() not in headers]
    if missing:
        raise SamplingError(f"{source} is missing required columns: {', '.join(missing)}")


def load_rule_mapping(path: str | Path) -> dict[str, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = _header_map(next(rows))
    except StopIteration as exc:
        raise SamplingError("The mapping workbook is empty.") from exc
    _require_headers(headers, ("SAM Rule ID", "Sampling Method"), "Mapping workbook")

    mapping: dict[str, str] = {}
    for row_number, row in enumerate(rows, start=2):
        rule_id = _text(row[headers["sam rule id"]])
        method = _text(row[headers["sampling method"]])
        if not rule_id and not method:
            continue
        if not rule_id or not method:
            raise SamplingError(f"Incomplete rule mapping on row {row_number}.")
        if rule_id in mapping and mapping[rule_id].casefold() != method.casefold():
            raise SamplingError(f"Rule ID {rule_id!r} has conflicting sampling methods.")
        mapping[rule_id] = method
    return mapping


def _decimal(value: object, row_number: int) -> Decimal:
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError) as exc:
        raise SamplingError(f"Invalid Amount on transaction row {row_number}: {value!r}") from exc


def load_transactions(path: str | Path) -> list[Transaction]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = _header_map(next(rows))
    except StopIteration as exc:
        raise SamplingError("The transaction workbook is empty.") from exc

    required = (
        "Alert ID",
        "Rule ID",
        "Amount",
        "Orig Party",
        "Originator Account Number",
        "Originator Name",
        "Bene Party",
        "Beneficiary Account Number",
        "Beneficiary Name",
    )
    _require_headers(headers, required, "Transaction workbook")

    transactions: list[Transaction] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None for value in row):
            continue
        alert_id = _text(row[headers["alert id"]])
        rule_id = _text(row[headers["rule id"]])
        if not alert_id or not rule_id:
            raise SamplingError(f"Transaction row {row_number} has no Alert ID or Rule ID.")
        transactions.append(
            Transaction(
                source_row=row_number,
                alert_id=alert_id,
                rule_id=rule_id,
                amount=_decimal(row[headers["amount"]], row_number),
                originator_id=_text(row[headers["orig party"]]),
                originator_account=_text(row[headers["originator account number"]]),
                originator_name=_text(row[headers["originator name"]]),
                beneficiary_id=_text(row[headers["bene party"]]),
                beneficiary_account=_text(row[headers["beneficiary account number"]]),
                beneficiary_name=_text(row[headers["beneficiary name"]]),
            )
        )
    if not transactions:
        raise SamplingError("The transaction workbook contains no data rows.")
    return transactions


def validate_instructions(instructions: str) -> None:
    normalized = " ".join(instructions.casefold().replace("&", "and").split())
    if not normalized:
        raise SamplingError("Sampling instructions are empty.")
    has_method = "top 3 and 3" in normalized or "top three and three" in normalized
    required_concepts = ("cumulative", "volume")
    missing = [concept for concept in required_concepts if concept not in normalized]
    if not has_method or missing:
        details = "Top 3 & 3"
        if missing:
            details += f" with concepts: {', '.join(missing)}"
        raise SamplingError(f"Instructions do not define the supported methodology: {details}.")


def run_sampling(
    instructions: str,
    mapping_path: str | Path,
    transaction_path: str | Path,
    focal_party_overrides: dict[str, str] | None = None,
) -> list[AlertSamplingResult]:
    validate_instructions(instructions)
    mapping = load_rule_mapping(mapping_path)
    transactions = load_transactions(transaction_path)
    grouped: dict[str, list[Transaction]] = defaultdict(list)
    for transaction in transactions:
        grouped[transaction.alert_id].append(transaction)

    results: list[AlertSamplingResult] = []
    for alert_id in sorted(grouped):
        alert_transactions = grouped[alert_id]
        missing_rules = sorted(
            {item.rule_id for item in alert_transactions if item.rule_id not in mapping}
        )
        if missing_rules:
            raise SamplingError(
                f"Alert {alert_id} contains Rule IDs absent from the mapping: "
                f"{', '.join(missing_rules)}"
            )
        methods = {mapping[item.rule_id] for item in alert_transactions}
        if len({method.casefold() for method in methods}) != 1:
            raise SamplingError(
                f"Alert {alert_id} resolves to multiple sampling methods: "
                f"{', '.join(sorted(methods))}"
            )
        method = next(iter(methods))
        override = (focal_party_overrides or {}).get(alert_id)
        results.append(sample_alert(alert_transactions, method, override))
    return results


def _write_table(worksheet, headers: list[str], rows: Iterable[Iterable[object]]) -> None:
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, start=1):
        width = min(max(len(_text(cell.value)) for cell in column) + 2, 60)
        worksheet.column_dimensions[get_column_letter(index)].width = max(width, 12)


def write_results(
    output_path: str | Path,
    results: list[AlertSamplingResult],
    instructions: str,
    mapping_path: str | Path,
    transaction_path: str | Path,
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Run Summary"
    metadata = [
        ("Generated UTC", datetime.now(timezone.utc).isoformat()),
        ("Instructions SHA-256", sha256(instructions.encode("utf-8")).hexdigest()),
        ("Mapping file", str(Path(mapping_path).resolve())),
        ("Transaction file", str(Path(transaction_path).resolve())),
        ("Alerts processed", len(results)),
        ("Counterparties selected", sum(len(result.selected) for result in results)),
    ]
    _write_table(summary_sheet, ["Field", "Value"], metadata)

    selected_sheet = workbook.create_sheet("Sampled Counterparties")
    population_sheet = workbook.create_sheet("Counterparty Population")
    headers = [
        "Alert ID",
        "Sampling Method",
        "Counterparty Key",
        "Counterparty ID",
        "Counterparty Name",
        "Cumulative Value",
        "Transaction Volume",
        "Value Rank",
        "Volume Rank",
        "Selected By Value",
        "Selected By Volume",
        "Selection Reason",
        "Transaction Sides",
        "Source Rows",
        "Rule IDs",
        "Focal Party ID",
        "Focal Party Name",
        "Focal Party Key",
        "Focal Party Source",
    ]

    def result_rows(selected_only: bool):
        for result in results:
            counterparties = result.selected if selected_only else result.population
            for counterparty in counterparties:
                yield (
                    result.alert_id,
                    result.sampling_method,
                    counterparty.key,
                    counterparty.party_id,
                    counterparty.name,
                    float(counterparty.cumulative_value),
                    counterparty.transaction_volume,
                    counterparty.value_rank,
                    counterparty.volume_rank,
                    "Yes" if counterparty.selected_by_value else "No",
                    "Yes" if counterparty.selected_by_volume else "No",
                    counterparty.selection_reason,
                    ", ".join(counterparty.sides),
                    ", ".join(str(row) for row in counterparty.source_rows),
                    ", ".join(result.rule_ids),
                    result.focal_party_id,
                    result.focal_party_name,
                    result.focal_party_key,
                    result.focal_inference,
                )

    _write_table(selected_sheet, headers, result_rows(selected_only=True))
    _write_table(population_sheet, headers, result_rows(selected_only=False))
    workbook.save(output_path)
