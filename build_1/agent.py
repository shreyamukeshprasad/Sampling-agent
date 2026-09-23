"""Rule-level counterparty sampling for alerted transaction workbooks."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import re
import sys
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class SamplingError(ValueError):
    """The input cannot be sampled reliably."""


@dataclass(frozen=True)
class Party:
    key: str
    party_id: str
    name: str
    side: str


@dataclass(frozen=True)
class Transaction:
    source_row: int
    alert_id: str
    rule_id: str
    amount: Decimal
    originator_id: str = ""
    originator_account: str = ""
    originator_name: str = ""
    beneficiary_id: str = ""
    beneficiary_account: str = ""
    beneficiary_name: str = ""

    def parties(self) -> tuple[Party | None, Party | None]:
        return (
            make_party(
                self.originator_id,
                self.originator_account,
                self.originator_name,
                "Originator",
            ),
            make_party(
                self.beneficiary_id,
                self.beneficiary_account,
                self.beneficiary_name,
                "Beneficiary",
            ),
        )


@dataclass
class CounterpartyCalculation:
    key: str
    party_id: str
    name: str
    cumulative_value: Decimal
    transaction_volume: int
    sides: tuple[str, ...]
    source_rows: tuple[int, ...]
    value_rank: int = 0
    volume_rank: int = 0
    selected_by_value: bool = False
    selected_by_volume: bool = False
    selected_as_full_population: bool = False

    @property
    def selected(self) -> bool:
        return (
            self.selected_by_value
            or self.selected_by_volume
            or self.selected_as_full_population
        )

    @property
    def selection_reason(self) -> str:
        reasons: list[str] = []
        if self.selected_as_full_population:
            reasons.append("Entire population has fewer than 3 counterparties")
        if self.selected_by_value:
            reasons.append("Top 3 by cumulative value")
        if self.selected_by_volume:
            reasons.append("Top 3 by transaction volume")
        return "; ".join(reasons)


@dataclass
class RuleSample:
    alert_id: str
    rule_id: str
    sampling_method: str
    focal_party_key: str
    focal_party_id: str
    focal_party_name: str
    focal_party_source: str
    calculations: list[CounterpartyCalculation]

    @property
    def sampled_counterparties(self) -> list[CounterpartyCalculation]:
        return [item for item in self.calculations if item.selected]


@dataclass
class _Totals:
    value: Decimal = Decimal("0")
    source_rows: set[int] = field(default_factory=set)
    sides: set[str] = field(default_factory=set)
    party_ids: Counter[str] = field(default_factory=Counter)
    names: Counter[str] = field(default_factory=Counter)


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", clean_text(value).upper())


def make_party(party_id: object, account: object, name: object, side: str) -> Party | None:
    clean_id = clean_text(party_id)
    clean_account = clean_text(account)
    clean_name = clean_text(name)
    if clean_id:
        key = f"ID:{normalize(clean_id)}"
    elif clean_account:
        key = f"ACCOUNT:{normalize(clean_account)}"
    elif clean_name:
        key = f"NAME:{normalize(clean_name)}"
    else:
        return None
    return Party(key=key, party_id=clean_id, name=clean_name, side=side)


def representative(values: Counter[str]) -> str:
    if not values:
        return ""
    return sorted(values.items(), key=lambda item: (-item[1], item[0].upper()))[0][0]


def focal_side_from_rule(rule_id: str) -> str | None:
    tokens = {token for token in re.split(r"[^A-Z0-9]+", rule_id.upper()) if token}
    if tokens.intersection({"BEN", "BENE", "BENEFICIARY"}):
        return "Beneficiary"
    if tokens.intersection({"ORG", "ORI", "ORIG", "ORIGINATOR"}):
        return "Originator"
    return None


def infer_focal_party(transactions: list[Transaction]) -> tuple[str, str]:
    directional_counts: Counter[str] = Counter()
    total_counts: Counter[str] = Counter()
    originator_keys: set[str] = set()
    beneficiary_keys: set[str] = set()

    for transaction in transactions:
        originator, beneficiary = transaction.parties()
        if originator:
            originator_keys.add(originator.key)
            total_counts[originator.key] += 1
        if beneficiary:
            beneficiary_keys.add(beneficiary.key)
            total_counts[beneficiary.key] += 1

        focal_side = focal_side_from_rule(transaction.rule_id)
        focal_party = originator if focal_side == "Originator" else beneficiary
        if focal_side and focal_party:
            directional_counts[focal_party.key] += 1

    if directional_counts:
        candidates = directional_counts.most_common()
        if len(candidates) == 1 or candidates[0][1] > candidates[1][1]:
            return candidates[0][0], "Inferred from Rule ID direction and party frequency"

    candidates_on_both_sides = originator_keys.intersection(beneficiary_keys)
    if candidates_on_both_sides:
        candidates = sorted(
            candidates_on_both_sides,
            key=lambda key: (-total_counts[key], key),
        )
        if len(candidates) == 1 or total_counts[candidates[0]] > total_counts[candidates[1]]:
            return candidates[0], "Inferred as the recurring party on both transaction sides"

    raise SamplingError(
        "Could not infer one focal party. Provide an Alert ID and Rule ID focal override."
    )


def resolve_focal_override(transactions: list[Transaction], party_id: str) -> str:
    requested = normalize(party_id)
    matches = {
        party.key
        for transaction in transactions
        for party in transaction.parties()
        if party and normalize(party.party_id) == requested
    }
    if not matches:
        raise SamplingError(f"Focal party override {party_id!r} was not found.")
    if len(matches) > 1:
        raise SamplingError(f"Focal party override {party_id!r} matched multiple parties.")
    return matches.pop()


def calculate_counterparties(
    transactions: list[Transaction], focal_party_key: str
) -> tuple[list[CounterpartyCalculation], str, str]:
    totals: dict[str, _Totals] = defaultdict(_Totals)
    focal_ids: Counter[str] = Counter()
    focal_names: Counter[str] = Counter()

    for transaction in transactions:
        parties = [party for party in transaction.parties() if party]
        for party in parties:
            if party.key == focal_party_key:
                if party.party_id:
                    focal_ids[party.party_id] += 1
                if party.name:
                    focal_names[party.name] += 1

        # Count a transaction once for each distinct non-focal party on that row.
        counterparties = {party.key: party for party in parties if party.key != focal_party_key}
        for party in counterparties.values():
            party_totals = totals[party.key]
            party_totals.value += abs(transaction.amount)
            party_totals.source_rows.add(transaction.source_row)
            party_totals.sides.add(party.side)
            if party.party_id:
                party_totals.party_ids[party.party_id] += 1
            if party.name:
                party_totals.names[party.name] += 1

    calculations = [
        CounterpartyCalculation(
            key=key,
            party_id=representative(item.party_ids),
            name=representative(item.names),
            cumulative_value=item.value,
            transaction_volume=len(item.source_rows),
            sides=tuple(sorted(item.sides)),
            source_rows=tuple(sorted(item.source_rows)),
        )
        for key, item in totals.items()
    ]
    return calculations, representative(focal_ids), representative(focal_names)


def apply_top_three_and_three(
    calculations: list[CounterpartyCalculation],
) -> list[CounterpartyCalculation]:
    by_value = sorted(
        calculations,
        key=lambda item: (-item.cumulative_value, -item.transaction_volume, item.key),
    )
    by_volume = sorted(
        calculations,
        key=lambda item: (-item.transaction_volume, -item.cumulative_value, item.key),
    )

    for rank, counterparty in enumerate(by_value, start=1):
        counterparty.value_rank = rank
    for rank, counterparty in enumerate(by_volume, start=1):
        counterparty.volume_rank = rank

    if len(calculations) < 3:
        for counterparty in calculations:
            counterparty.selected_as_full_population = True
    else:
        for counterparty in by_value[:3]:
            counterparty.selected_by_value = True
        for counterparty in by_volume[:3]:
            counterparty.selected_by_volume = True

    return sorted(calculations, key=lambda item: (not item.selected, item.value_rank, item.key))


def sample_one_rule(
    transactions: Iterable[Transaction],
    sampling_method: str,
    focal_party_id: str | None = None,
) -> RuleSample:
    rows = list(transactions)
    if not rows:
        raise SamplingError("A rule sample must contain at least one transaction.")
    groups = {(row.alert_id, row.rule_id) for row in rows}
    if len(groups) != 1:
        raise SamplingError("Each calculation must contain one Alert ID and one Rule ID.")
    if sampling_method.strip().casefold() != "top 3 & 3":
        raise SamplingError(f"Unsupported sampling method: {sampling_method}")

    if focal_party_id:
        focal_key = resolve_focal_override(rows, focal_party_id)
        focal_source = "Provided explicitly"
    else:
        focal_key, focal_source = infer_focal_party(rows)

    calculations, focal_id, focal_name = calculate_counterparties(rows, focal_key)
    calculations = apply_top_three_and_three(calculations)
    alert_id, rule_id = groups.pop()
    return RuleSample(
        alert_id=alert_id,
        rule_id=rule_id,
        sampling_method="Top 3 & 3",
        focal_party_key=focal_key,
        focal_party_id=focal_id,
        focal_party_name=focal_name,
        focal_party_source=focal_source,
        calculations=calculations,
    )


def header_positions(headers: Iterable[object]) -> dict[str, int]:
    return {clean_text(header).casefold(): index for index, header in enumerate(headers)}


def require_columns(columns: dict[str, int], required: Iterable[str], source: str) -> None:
    missing = [column for column in required if column.casefold() not in columns]
    if missing:
        raise SamplingError(f"{source} is missing columns: {', '.join(missing)}")


def load_rule_mapping(path: Path) -> dict[str, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        columns = header_positions(next(rows))
    except StopIteration as exc:
        raise SamplingError("The mapping workbook is empty.") from exc
    require_columns(columns, ("SAM Rule ID", "Sampling Method"), "Mapping workbook")

    mapping: dict[str, str] = {}
    for row_number, row in enumerate(rows, start=2):
        rule_id = clean_text(row[columns["sam rule id"]])
        method = clean_text(row[columns["sampling method"]])
        if not rule_id and not method:
            continue
        if not rule_id or not method:
            raise SamplingError(f"Incomplete mapping on row {row_number}.")
        if rule_id in mapping and mapping[rule_id].casefold() != method.casefold():
            raise SamplingError(f"Rule ID {rule_id!r} has conflicting sampling methods.")
        mapping[rule_id] = method
    return mapping


def parse_amount(value: object, row_number: int) -> Decimal:
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError) as exc:
        raise SamplingError(f"Invalid Amount on transaction row {row_number}: {value!r}") from exc


def load_transactions(path: Path) -> list[Transaction]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        columns = header_positions(next(rows))
    except StopIteration as exc:
        raise SamplingError("The transaction workbook is empty.") from exc

    required = (
        "Alert ID",
        "Rule ID",
        "Amount",
        "Orig Party",
        "Originator Account Number",
        "Originator Name",
        "Bene Party",
        "Beneficiary Account Number",
        "Beneficiary Name",
    )
    require_columns(columns, required, "Transaction workbook")

    transactions: list[Transaction] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None for value in row):
            continue
        alert_id = clean_text(row[columns["alert id"]])
        rule_id = clean_text(row[columns["rule id"]])
        if not alert_id or not rule_id:
            raise SamplingError(f"Transaction row {row_number} has no Alert ID or Rule ID.")
        transactions.append(
            Transaction(
                source_row=row_number,
                alert_id=alert_id,
                rule_id=rule_id,
                amount=parse_amount(row[columns["amount"]], row_number),
                originator_id=clean_text(row[columns["orig party"]]),
                originator_account=clean_text(row[columns["originator account number"]]),
                originator_name=clean_text(row[columns["originator name"]]),
                beneficiary_id=clean_text(row[columns["bene party"]]),
                beneficiary_account=clean_text(row[columns["beneficiary account number"]]),
                beneficiary_name=clean_text(row[columns["beneficiary name"]]),
            )
        )
    if not transactions:
        raise SamplingError("The transaction workbook contains no data rows.")
    return transactions


def validate_prompt(prompt: str) -> None:
    normalized_prompt = " ".join(prompt.casefold().replace("&", "and").split())
    required_terms = {
        "Top 3 & 3": "top 3 and 3",
        "rule-level sampling": "rule id",
        "cumulative value": "cumulative value",
        "transaction volume": "transaction volume",
        "focal-party exclusion": "focal",
    }
    missing = [label for label, phrase in required_terms.items() if phrase not in normalized_prompt]
    if missing:
        raise SamplingError(f"Sampling prompt is missing: {', '.join(missing)}")


def sample_all_rules(
    prompt: str,
    mapping_path: Path,
    transaction_path: Path,
    focal_overrides: dict[tuple[str, str], str] | None = None,
) -> list[RuleSample]:
    mapping = load_rule_mapping(mapping_path)
    transactions = load_transactions(transaction_path)
    return sample_transaction_groups(prompt, mapping, transactions, focal_overrides)


def sample_transaction_groups(
    prompt: str,
    mapping: dict[str, str],
    transactions: Iterable[Transaction],
    focal_overrides: dict[tuple[str, str], str] | None = None,
) -> list[RuleSample]:
    validate_prompt(prompt)
    grouped: dict[tuple[str, str], list[Transaction]] = defaultdict(list)
    for transaction in transactions:
        grouped[(transaction.alert_id, transaction.rule_id)].append(transaction)

    samples: list[RuleSample] = []
    for group_key in sorted(grouped):
        alert_id, rule_id = group_key
        if rule_id not in mapping:
            raise SamplingError(f"Rule ID {rule_id!r} is absent from the mapping workbook.")
        focal_override = (focal_overrides or {}).get(group_key)
        try:
            sample = sample_one_rule(grouped[group_key], mapping[rule_id], focal_override)
        except SamplingError as exc:
            raise SamplingError(f"Alert {alert_id}, Rule {rule_id}: {exc}") from exc
        samples.append(sample)
    return samples


RESULT_HEADERS = [
    "Alert ID",
    "Rule ID",
    "Sampling Method",
    "Counterparty Key",
    "Counterparty ID",
    "Counterparty Name",
    "Cumulative Value",
    "Transaction Volume",
    "Value Rank",
    "Volume Rank",
    "Selected By Value",
    "Selected By Volume",
    "Selection Reason",
    "Transaction Sides",
    "Source Rows",
    "Focal Party ID",
    "Focal Party Name",
    "Focal Party Source",
]


def result_row(sample: RuleSample, item: CounterpartyCalculation) -> tuple[object, ...]:
    return (
        sample.alert_id,
        sample.rule_id,
        sample.sampling_method,
        item.key,
        item.party_id,
        item.name,
        float(item.cumulative_value),
        item.transaction_volume,
        item.value_rank,
        item.volume_rank,
        "Yes" if item.selected_by_value else "No",
        "Yes" if item.selected_by_volume else "No",
        item.selection_reason,
        ", ".join(item.sides),
        ", ".join(str(row) for row in item.source_rows),
        sample.focal_party_id,
        sample.focal_party_name,
        sample.focal_party_source,
    )


def format_table(worksheet, header_row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    for index, column in enumerate(worksheet.columns, start=1):
        width = min(max(len(clean_text(cell.value)) for cell in column) + 2, 55)
        worksheet.column_dimensions[get_column_letter(index)].width = max(width, 12)
    for row in worksheet.iter_rows(min_row=header_row + 1):
        row[6].number_format = '#,##0.00'


def add_table_sheet(workbook: Workbook, name: str, rows: Iterable[tuple[object, ...]]) -> None:
    worksheet = workbook.create_sheet(name)
    worksheet.append(RESULT_HEADERS)
    for row in rows:
        worksheet.append(row)
    format_table(worksheet)


def write_results(
    output_path: Path,
    samples: list[RuleSample],
    prompt: str,
    mapping_path: Path,
    transaction_path: Path,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Run Summary"
    summary_rows = [
        ("Generated UTC", datetime.now(timezone.utc).isoformat()),
        ("Sampling prompt SHA-256", sha256(prompt.encode("utf-8")).hexdigest()),
        ("Mapping file", str(mapping_path.resolve())),
        ("Transaction file", str(transaction_path.resolve())),
        ("Rule samples created", len(samples)),
        ("Sampled counterparties", sum(len(sample.sampled_counterparties) for sample in samples)),
    ]
    summary.append(("Field", "Value"))
    for row in summary_rows:
        summary.append(row)
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 90

    add_table_sheet(
        workbook,
        "Sampled Counterparties",
        (
            result_row(sample, item)
            for sample in samples
            for item in sample.sampled_counterparties
        ),
    )
    add_table_sheet(
        workbook,
        "All Calculations",
        (result_row(sample, item) for sample in samples for item in sample.calculations),
    )

    for index, sample in enumerate(samples, start=1):
        worksheet = workbook.create_sheet(f"Rule {index}")
        worksheet.append(("Alert ID", sample.alert_id))
        worksheet.append(("Rule ID", sample.rule_id))
        worksheet.append(("Sampling Method", sample.sampling_method))
        worksheet.append(("Focal Party", sample.focal_party_id or sample.focal_party_key))
        worksheet.append(("Focal Party Source", sample.focal_party_source))
        worksheet.append(())
        worksheet.append(RESULT_HEADERS)
        for item in sample.sampled_counterparties:
            worksheet.append(result_row(sample, item))
        for cell in worksheet[7]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        worksheet.freeze_panes = "A8"
        worksheet.auto_filter.ref = (
            f"A7:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
        )
        for column_index, column in enumerate(worksheet.columns, start=1):
            width = min(max(len(clean_text(cell.value)) for cell in column) + 2, 55)
            worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)
        for row in worksheet.iter_rows(min_row=8):
            row[6].number_format = '#,##0.00'

    workbook.save(output_path)


def parse_override(value: str) -> tuple[tuple[str, str], str]:
    if "=" not in value or "|" not in value.split("=", 1)[0]:
        raise argparse.ArgumentTypeError(
            "Use ALERT_ID|RULE_ID=PARTY_ID for a focal-party override."
        )
    group, party_id = value.split("=", 1)
    alert_id, rule_id = group.split("|", 1)
    if not alert_id.strip() or not rule_id.strip() or not party_id.strip():
        raise argparse.ArgumentTypeError(
            "Use ALERT_ID|RULE_ID=PARTY_ID for a focal-party override."
        )
    return (alert_id.strip(), rule_id.strip()), party_id.strip()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Sample counterparties separately for each rule.")
    parser.add_argument("--prompt", required=True, type=Path)
    parser.add_argument("--mapping", required=True, type=Path)
    parser.add_argument("--transactions", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument(
        "--focal-party",
        action="append",
        default=[],
        type=parse_override,
        metavar="ALERT_ID|RULE_ID=PARTY_ID",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        prompt = args.prompt.read_text(encoding="utf-8")
        samples = sample_all_rules(
            prompt,
            args.mapping,
            args.transactions,
            focal_overrides=dict(args.focal_party),
        )
        args.output.parent.mkdir(parents=True, exist_ok=True)
        write_results(args.output, samples, prompt, args.mapping, args.transactions)
    except (OSError, SamplingError) as exc:
        print(f"Sampling failed: {exc}", file=sys.stderr)
        return 1

    print(f"Created {len(samples)} rule sample(s): {args.output.resolve()}")
    for sample in samples:
        print(
            f"- {sample.rule_id}: {len(sample.sampled_counterparties)} sampled "
            f"from {len(sample.calculations)} calculated counterparties"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
